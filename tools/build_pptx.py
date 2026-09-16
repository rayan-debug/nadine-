#!/usr/bin/env python3
"""Build the editable PowerPoint deck from the measured layout manifest.

Every title, statement, list, reference, evidence chip and page number arrives
as a real text box in Cormorant Garamond / Montserrat; every rule, panel,
pyramid tier, axis and pill is a drawn shape; only the bespoke anatomical
illustrations are placed as high-resolution images, and their editable .svg
sources travel in deck/illustrations/.

Presenter notes are written into the notes pane of each slide.
"""

import json
import pathlib
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = pathlib.Path(__file__).resolve().parent.parent
DECK = ROOT / "deck"
BUILD = DECK / "build"
IMG = BUILD / "img"
OUT = DECK / "Tissue-Specific-Intimate-Medicine.pptx"

# the stage is authored at 1920 x 1080; a 16:9 slide is 13.333 x 7.5 in
SLIDE_W, SLIDE_H = Inches(13.3333), Inches(7.5)
PXE = SLIDE_W / 1920                       # EMU per authored pixel
PAPER = RGBColor(0xFB, 0xF7, 0xEE)

ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
         "right": PP_ALIGN.RIGHT, "start": PP_ALIGN.LEFT, "end": PP_ALIGN.RIGHT,
         "justify": PP_ALIGN.JUSTIFY}


def E(px):
    return Emu(int(round(px * PXE)))


def C(rgb):
    return RGBColor(*rgb)


def set_tracking(run, px):
    """Letter-spacing, which python-pptx does not expose."""
    if not px:
        return
    run.font._rPr.set("spc", str(int(round(px * 50))))   # px -> pt -> 1/100 pt


def no_autofit(tf):
    # spAutoFit plus wrap="none" makes LibreOffice shrink the box to the text
    # and re-centre it, which silently re-aligns every left-set heading
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0


def add_text(slide, t):
    b, pad = t["box"], t["pad"]
    w = max(b["w"] - pad[1] - pad[3], 1)
    h = max(b["h"] - pad[0] - pad[2], 1)
    hard_breaks = sum(r["t"] == "\n" for r in t["runs"])
    # PowerPoint measures type slightly differently from the browser, so a line
    # the browser fitted can wrap and push the slide apart. Anything that was a
    # single line is therefore pinned as a single line.
    single = round(h / t["line"]) <= 1 and not hard_breaks

    box = slide.shapes.add_textbox(
        E(b["x"] + pad[3] - (0 if single else 5)), E(b["y"] + pad[0]),
        E(w + (0 if single else 10)), E(h + (0 if single else t["line"])))
    tf = box.text_frame
    no_autofit(tf)
    tf.word_wrap = not single
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE if single else MSO_ANCHOR.TOP
    font = "Cormorant Garamond" if t["serif"] else "Montserrat"

    paras = [[]]
    for r in t["runs"]:
        parts = r["t"].split("\n")
        for k, part in enumerate(parts):
            if k:
                paras.append([])
            if part:
                paras[-1].append({**r, "t": part})

    for pi, runs in enumerate(paras):
        if runs:
            runs[0]["t"] = runs[0]["t"].lstrip()
            runs[-1]["t"] = runs[-1]["t"].rstrip()
        p = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
        p.alignment = ALIGN.get(t["align"], PP_ALIGN.LEFT)
        p.line_spacing = Pt(t["line"] / 2)
        for r in runs:
            run = p.add_run()
            run.text = r["t"].upper() if t["caps"] else r["t"]
            f = run.font
            f.name = font
            f.size = Pt(t["size"] / 2)
            f.bold = bool(r["b"])
            f.italic = t["italic"]
            f.color.rgb = C(r["c"])
            set_tracking(run, t["track"])
    return box


def line(slide, x1, y1, x2, y2, colour, w):
    from pptx.util import Emu as _E
    ln = slide.shapes.add_connector(1, E(x1), E(y1), E(x2), E(y2))  # 1 = straight
    ln.line.color.rgb = C(colour)
    ln.line.width = _E(max(int(round(w * PXE)), 6350))
    return ln


def add_shape(slide, s):
    b, br = s["box"], s["borders"]

    if s["clip"] and s["clip"].startswith("polygon"):
        # Chromium reports the resolved polygon in a mix of px and %
        vals = re.findall(r"(-?[\d.]+)(px|%)", s["clip"])
        if len(vals) >= 6:
            co = [float(v) * (b["w" if i % 2 == 0 else "h"] / 100 if u == "%" else 1)
                  for i, (v, u) in enumerate(vals)]
            pts = [(b["x"] + co[i], b["y"] + co[i + 1]) for i in range(0, len(co) - 1, 2)]
            fb = slide.shapes.build_freeform(E(pts[0][0]), E(pts[0][1]))
            fb.add_line_segments([(E(x), E(y)) for x, y in pts[1:]], close=True)
            sh = fb.convert_to_shape()
            sh.line.fill.background()
            sh.shadow.inherit = False
            if s["fill"]:
                sh.fill.solid()
                sh.fill.fore_color.rgb = C(s["fill"])
            else:
                sh.fill.background()
            return sh

    # the evidence axis is painted with a CSS gradient, so it carries no
    # background-color to copy - draw it as a gold rule instead
    if "axis" in s["sel"]:
        return line(slide, b["x"], b["y"] + b["h"] / 2, b["x"] + b["w"],
                    b["y"] + b["h"] / 2, (0xC6, 0xA2, 0x4E), max(b["h"], 2))

    uniform = all(x["w"] == br[0]["w"] and x["c"] == br[0]["c"] for x in br)

    if s["fill"] or (uniform and br[0]["w"]):
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if s["radius"] > 2 else MSO_SHAPE.RECTANGLE
        if s["sel"].split()[0] == "pin":
            kind = MSO_SHAPE.OVAL
        sh = slide.shapes.add_shape(kind, E(b["x"]), E(b["y"]), E(b["w"]), E(b["h"]))
        if s["radius"] > 2 and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            sh.adjustments[0] = min(0.5, s["radius"] / min(b["w"], b["h"]))
        if s["fill"]:
            sh.fill.solid()
            sh.fill.fore_color.rgb = C(s["fill"])
        else:
            sh.fill.background()
        if uniform and br[0]["w"]:
            sh.line.color.rgb = C(br[0]["c"])
            sh.line.width = Emu(max(int(round(br[0]["w"] * PXE)), 6350))
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        sh = _drop_placeholder_text(sh)

    # per-edge rules (footer hairline, cascade left border, take-home underline)
    edges = [(0, b["x"], b["y"], b["x"] + b["w"], b["y"]),
             (1, b["x"] + b["w"], b["y"], b["x"] + b["w"], b["y"] + b["h"]),
             (2, b["x"], b["y"] + b["h"], b["x"] + b["w"], b["y"] + b["h"]),
             (3, b["x"], b["y"], b["x"], b["y"] + b["h"])]
    if not uniform:
        for i, x1, y1, x2, y2 in edges:
            if br[i]["w"]:
                line(slide, x1, y1, x2, y2, br[i]["c"], br[i]["w"])
    return None


def _drop_placeholder_text(sh):
    if sh.has_text_frame:
        sh.text_frame.word_wrap = False
    return sh


def main():
    manifest = json.loads((BUILD / "layout.json").read_text())
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    blank = prs.slide_layouts[6]

    for m in manifest:
        slide = prs.slides.add_slide(blank)
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = PAPER

        for s in m["shapes"]:
            add_shape(slide, s)
        for im in m["images"]:
            b = im["box"]
            slide.shapes.add_picture(str(IMG / im["file"]),
                                     E(b["x"]), E(b["y"]), E(b["w"]), E(b["h"]))
        for t in m["texts"]:
            add_text(slide, t)

        tf = slide.notes_slide.notes_text_frame
        tf.text = m["notes"][0] if m["notes"] else ""
        for para in m["notes"][1:]:
            tf.add_paragraph().text = para

    prs.save(OUT)
    kb = OUT.stat().st_size / 1024
    print(f"  {OUT.name}  {kb:.0f} kB  ({len(manifest)} slides, notes preserved)")


if __name__ == "__main__":
    print("pptx ->", DECK)
    main()
